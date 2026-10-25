# -*- coding: utf-8 -*-
"""
Script to generate the complete EVATRACK Software Requirements Specification (SRS) v1.0
following the official University of San Jose-Recoletos (USJ-R) Capstone guidelines,
incorporating all Must-Haves from EVATRACK-Must-Haves (1).docx and the production codebase.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background shading of a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding/margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    """Applies clean, subtle borders to an entire table."""
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

def style_table(table, col_widths=None, header_bg="1E3A8A", alt_row=True):
    """Applies styling, headers, widths, and zebra-striping to tables."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Style Header Row
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.name = "Calibri"

    # Style Data Rows
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg = "F8FAFC" if (alt_row and row_idx % 2 == 1) else "FFFFFF"
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.size = Pt(9.0)
                    r.font.name = "Calibri"
                    r.font.color.rgb = RGBColor(30, 41, 59)

    # Set Column Widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

def add_styled_heading(doc, text, level):
    """Adds standard formatted headings with consistent color and spacing."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.bold = True

    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(30, 58, 138) # Navy Blue
    elif level == 2:
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=6):
    """Adds a standard body paragraph with 1.15 line spacing."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(doc, text, bold_prefix=None):
    """Adds a formatted bullet item."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def build_srs_document(output_path):
    doc = docx.Document()

    # Set 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)

    # =========================================================================
    # COVER PAGE
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("UNIVERSITY OF SAN JOSE-RECOLETOS")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("School of Computer Studies")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(71, 85, 105)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(45)
    r = p.add_run("Cebu City, Philippines")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(100, 116, 139)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("EVATRACK")
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Streamlined Evacuation Processes and Coordination Platform Integrating SafeTrack Demographics and ResQperation Response Support")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Software Requirements Specification")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(71, 85, 105)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(40)
    r = p.add_run("Version 1.0")
    r.font.name = "Calibri"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    # Capstone Dedication
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("A Capstone Project Presented to the Faculty of the\nSchool of Computer Studies\nUniversity of San Jose-Recoletos")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(71, 85, 105)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("In Partial Fulfillment of the Requirements for\n")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(71, 85, 105)
    r_deg = p.add_run("Bachelor of Science in Information Technology / Associate in Computer Technology")
    r_deg.font.name = "Calibri"
    r_deg.font.size = Pt(10.5)
    r_deg.font.bold = True
    r_deg.font.color.rgb = RGBColor(15, 23, 42)

    # Proponents Box
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Presented by:")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)

    proponents = [
        "Klint M. Ruales",
        "Danica Gelbolingo",
        "Anna Rhea Villadolid",
        "Vhenz Cernal"
    ]
    for prop in proponents:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(prop)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(15, 23, 42)

    # Advisory Panel
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Capstone Adviser: ")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138)
    r_name = p.add_run("Mr. Stefan A.D. Flores")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Project Coordinator: ")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105)
    r_name = p.add_run("Mr. Roderick A. Bandalan")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(25)
    r = p.add_run("Panel Chair: ")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105)
    r_name = p.add_run("Mr. Eric Magto  |  ")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(10)
    r_pan = p.add_run("Panelist: ")
    r_pan.font.name = "Calibri"
    r_pan.font.size = Pt(10)
    r_pan.font.bold = True
    r_pan.font.color.rgb = RGBColor(71, 85, 105)
    r_pname = p.add_run("Ms. Khiara Danielle Rubia")
    r_pname.font.name = "Calibri"
    r_pname.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("October 2026")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # =========================================================================
    # REVISION HISTORY
    # =========================================================================
    add_styled_heading(doc, "Revision History", level=1)
    rev_table = doc.add_table(rows=4, cols=4)
    rev_headers = ["Version", "Date", "Description", "Author(s)"]
    for c_idx, h in enumerate(rev_headers):
        rev_table.cell(0, c_idx).text = h

    rev_data = [
        ("0.1", "2026-07-20", "Initial draft: project scope, must-haves elicitation, and stakeholder requirements analysis.", "All proponents"),
        ("0.5", "2026-08-20", "Complete functional and non-functional requirements, integration architecture (SafeTrack and ResQperation), adviser comments incorporated.", "Ruales, Gelbolingo, Villadolid, Cernal"),
        ("1.0", "2026-10-05", "Approved baseline for capstone panel defense incorporating verified backend actions, analytics computations, and field-tested UI workflows.", "All proponents")
    ]
    for r_idx, row in enumerate(rev_data, start=1):
        for c_idx, val in enumerate(row):
            rev_table.cell(r_idx, c_idx).text = val

    style_table(rev_table, col_widths=[Inches(0.8), Inches(1.1), Inches(3.2), Inches(1.4)])

    # =========================================================================
    # TABLE OF CONTENTS
    # =========================================================================
    add_styled_heading(doc, "Table of Contents", level=1)
    toc_items = [
        ("Revision History", "2"),
        ("Table of Contents", "3"),
        ("1. Introduction", "4"),
        ("    1.1 Purpose", "4"),
        ("    1.2 Scope", "4"),
        ("    1.3 Definitions, Acronyms, and Abbreviations", "5"),
        ("    1.4 References", "6"),
        ("    1.5 Overview", "6"),
        ("2. Overall Description", "7"),
        ("    2.1 Product Perspective", "7"),
        ("    2.2 Product Functions", "8"),
        ("    2.3 User Classes and Characteristics", "8"),
        ("    2.4 Operating Environment", "9"),
        ("    2.5 Design and Implementation Constraints", "9"),
        ("    2.6 Assumptions and Dependencies", "10"),
        ("    2.7 Apportioning of Requirements", "10"),
        ("3. Specific Requirements", "11"),
        ("    3.1 External Interface Requirements", "11"),
        ("        3.1.1 User Interfaces", "11"),
        ("        3.1.2 Hardware Interfaces", "11"),
        ("        3.1.3 Software Interfaces", "12"),
        ("        3.1.4 Communications Interfaces", "12"),
        ("    3.2 Functional Requirements", "13"),
        ("        3.2.1 Authentication & Role-Based Access Control (AUTH)", "13"),
        ("        3.2.2 SafeTrack Demographics Ingestion & Isolation (STK)", "14"),
        ("        3.2.3 Evacuation Center & Room Configuration (CTR)", "15"),
        ("        3.2.4 Household Verification & Admission Tracking (VER)", "16"),
        ("        3.2.5 In-Center Stay, Room Assignment & Checkout (STAY)", "17"),
        ("        3.2.6 Disaster Event Coordination (EVT)", "18"),
        ("        3.2.7 In-Center Needs & Incident Reporting (ISS)", "19"),
        ("        3.2.8 Resource Demands & ResQperation Response Routing (RES)", "20"),
        ("        3.2.9 Emergency Dispatches, SMS & Push Broadcasts (NOTIF)", "21"),
        ("        3.2.10 Evacuation Analytics & Demographic Reporting (ANL)", "22"),
        ("        3.2.11 Audit Logging & Super Admin System Oversight (AUD)", "23"),
        ("    3.3 Non-Functional Requirements", "24"),
        ("        3.3.1 Performance", "24"),
        ("        3.3.2 Security & Data Confidentiality", "24"),
        ("        3.3.3 Reliability and Disaster Availability", "25"),
        ("        3.3.4 Usability & Field Accessibility", "25"),
        ("        3.3.5 Maintainability & Clean Architecture", "26"),
        ("        3.3.6 Portability", "26"),
        ("        3.3.7 Data Privacy and Regulatory Compliance", "26"),
        ("    3.4 Database Requirements", "27"),
        ("Appendix A. Preliminary Entity-Relationship Overview", "30"),
        ("Appendix B. Requirements Traceability Matrix", "31"),
        ("Appendix C. Capstone Approval & Sign-Off Sheet", "34")
    ]
    for title, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(10)
        if not title.startswith("    "):
            r_t.font.bold = True
            r_t.font.color.rgb = RGBColor(30, 58, 138)
        else:
            r_t.font.color.rgb = RGBColor(51, 65, 85)

        # Tab to page
        r_dot = p.add_run(" " + "." * max(2, (75 - len(title) * 2)) + " ")
        r_dot.font.color.rgb = RGBColor(203, 213, 225)
        r_p = p.add_run(page)
        r_p.font.name = "Calibri"
        r_p.font.size = Pt(9.5)
        r_p.font.bold = True
        r_p.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # =========================================================================
    # 1. INTRODUCTION
    # =========================================================================
    add_styled_heading(doc, "1. Introduction", level=1)

    add_styled_heading(doc, "1.1 Purpose", level=2)
    add_body_p(doc, 
        "This Software Requirements Specification (SRS) establishes the formal functional, behavioral, and non-functional baseline for the development and evaluation of EVATRACK (Version 1.0) — a Streamlined Evacuation Processes and Coordination Platform. It defines the complete scope of system capabilities, constraints, external interfaces, and security controls agreed upon between the student proponents, the capstone adviser, the project coordinator, the evaluation panel of the School of Computer Studies at the University of San Jose-Recoletos (USJ-R), and emergency response stakeholders.",
        bold_prefix="Purpose and Objective: "
    )
    add_body_p(doc,
        "The intended audiences of this document are: (a) the Capstone Evaluation Panel and Project Coordinator, to assess architectural rigor, functional completeness, and compliance with college project standards; (b) the Capstone Adviser, to guide iterative technical execution; (c) the Project Development Team, to translate every verifiable requirement into production-ready software artifacts, unit/feature tests, and deployment scripts; and (d) Disaster Risk Reduction and Management Officers (DRRMO), Barangay Disaster Committees, and evacuation center personnel, to validate that the operational workflows accurately mitigate field challenges encountered during natural disaster evacuations."
    )

    add_styled_heading(doc, "1.2 Scope", level=2)
    add_body_p(doc,
        "EvaTrack is an enterprise-grade, web-based evacuation logistics, admission, and coordination management platform engineered to modernize disaster response within Local Government Units (LGUs). The system manages the full operational lifecycle of disaster evacuation across municipal shelters, comprising: multi-center registration and room-level capacity tracking; rapid evacuee household admission via QR code scanning and manual demographic search; in-center stay monitoring; emergency incident and facility health reporting; and resource demand management.",
        bold_prefix="System Scope: "
    )
    add_body_p(doc,
        "Crucially, EvaTrack establishes secure, bi-directional interoperability with two peer platforms in the disaster management ecosystem:\n"
        "(1) SafeTrack Demographics System: EvaTrack retrieves pre-registered household master records and vulnerability profiles from SafeTrack. In strict adherence to data privacy principles, this demographic data is completely hidden from public/resident view and is exposed solely to authorized evacuation personnel on-site for identity verification, demographic aggregation, and needs prioritization.\n"
        "(2) ResQperation Response Platform: EvaTrack routes resource supply requests, medical logistics packs, and search-and-rescue assistance directly to ResQperation coordinators, while utilizing ResQperation push gateways to broadcast real-time mobile evacuation warnings to citizens with smartphones."
    )
    add_body_p(doc,
        "The system enforces strict Role-Based Access Control (RBAC) across three primary operator tiers: Super Admin, Evacuation Admin, and Evacuation Personnel. Displaced households and citizen residents act as secondary beneficiaries who receive outbound multi-channel broadcasts (SMS and mobile push) and undergo on-site verification without holding direct interactive user accounts.",
        bold_prefix="User Roles & Access: "
    )
    add_body_p(doc,
        "Explicitly Out of Scope for Version 1.0:\n"
        "• Long-term permanent post-disaster housing re-settlement administration;\n"
        "• Direct financial cash-disbursement or commercial point-of-sale transactions;\n"
        "• Wearable biometric health sensor telemetry;\n"
        "• Native mobile app binary stores (EvaTrack provides a fully responsive PWA web interface accessed via desktop, tablet, and mobile browsers).",
        bold_prefix="Scope Exclusions: "
    )

    add_styled_heading(doc, "1.3 Definitions, Acronyms, and Abbreviations", level=2)
    def_table = doc.add_table(rows=17, cols=2)
    def_table.cell(0, 0).text = "Term / Acronym"
    def_table.cell(0, 1).text = "Definition"

    defs = [
        ("EvaTrack", "Streamlined Evacuation Processes and Coordination Platform specified in this document."),
        ("SafeTrack", "Municipal census and household demographic profiling system serving as the authoritative demographic data source."),
        ("ResQperation", "Emergency rescue dispatch and response coordination platform integrated with EvaTrack for logistics routing and push notifications."),
        ("SRS", "Software Requirements Specification — this document, conforming to IEEE Std 830-1998 and ISO/IEC/IEEE 29148:2018."),
        ("RBAC", "Role-Based Access Control — security architecture restricting system operations according to authenticated user roles."),
        ("LGU", "Local Government Unit (Municipality, City, or Barangay in the Republic of the Philippines)."),
        ("DRRMO / MDRRMO", "Disaster Risk Reduction and Management Office / Municipal DRRMO."),
        ("NDRRMC", "National Disaster Risk Reduction and Management Council of the Philippines."),
        ("DPA of 2012", "Republic Act No. 10173 (Data Privacy Act of 2012, Philippines) and its Implementing Rules and Regulations."),
        ("PWD", "Person with Disability — tracked as a vulnerable demographic category for priority evacuation accommodation."),
        ("QR Code", "Quick Response Matrix Barcode — encoded with a household's unique verification token for rapid contactless check-in."),
        ("SMS", "Short Message Service — utilized for dispatching cellular emergency advisories to residents without active mobile data."),
        ("REST API", "Representational State Transfer Application Programming Interface — JSON-based client-server communication architecture."),
        ("DTO", "Data Transfer Object — typed domain structure encapsulating payload parameters across architectural layers."),
        ("ORM", "Object-Relational Mapping — Eloquent abstraction layer mapping MySQL database tables to PHP domain models."),
        ("CRUD", "Create, Read, Update, Delete — fundamental persistent data operations executed by the backend repository layer.")
    ]
    for r_idx, (term, desc) in enumerate(defs, start=1):
        def_table.cell(r_idx, 0).text = term
        def_table.cell(r_idx, 1).text = desc

    style_table(def_table, col_widths=[Inches(1.8), Inches(4.7)])

    add_styled_heading(doc, "1.4 References", level=2)
    add_bullet(doc, "Systems and software engineering — Life cycle processes — Requirements engineering.", bold_prefix="ISO/IEC/IEEE 29148:2018: ")
    add_bullet(doc, "IEEE Recommended Practice for Software Requirements Specifications (structural reference).", bold_prefix="IEEE Std 830-1998: ")
    add_bullet(doc, "Data Privacy Act of 2012, National Privacy Commission, Republic of the Philippines.", bold_prefix="Republic Act No. 10173: ")
    add_bullet(doc, "Guidelines on Evacuation Center Coordination and Management, NDRRMC / DSWD / DOH.", bold_prefix="Joint Memorandum Circular No. 1 (2013): ")
    add_bullet(doc, "Laravel 11.x/12.x Official Documentation, Laravel LLC, https://laravel.com/docs.", bold_prefix="Laravel Framework Reference: ")
    add_bullet(doc, "React 18.x Documentation & Component API, Meta Platforms, Inc., https://react.dev.", bold_prefix="React Library Reference: ")
    add_bullet(doc, "MySQL 8.0 Reference Manual, Oracle Corporation, https://dev.mysql.com/doc/.", bold_prefix="MySQL 8.0 Engine Reference: ")
    add_bullet(doc, "EvaTrack Capstone Baseline Specification, School of Computer Studies, USJ-R.", bold_prefix="EVATRACK-Must-Haves (1).docx: ")

    add_styled_heading(doc, "1.5 Overview", level=2)
    add_body_p(doc,
        "The remainder of this document is structured as follows: Section 2 presents an overall description of EvaTrack, covering its operational context, high-level product functions, user classes, operating environment, design constraints, and dependencies. Section 3 details the specific requirements, including external interfaces (UI, hardware, software, communications), comprehensive functional requirements organized by domain subsystems, non-functional quality attributes, and the database schema dictionary. Appendices A, B, and C provide the preliminary Entity-Relationship narrative, the Must-Haves Requirements Traceability Matrix, and the Capstone Panel Approval Sheet."
    )

    doc.add_page_break()

    # =========================================================================
    # 2. OVERALL DESCRIPTION
    # =========================================================================
    add_styled_heading(doc, "2. Overall Description", level=1)

    add_styled_heading(doc, "2.1 Product Perspective", level=2)
    add_body_p(doc,
        "EvaTrack operates as a modern, decoupled web application built on Domain-Driven Design (DDD) principles. It bridges municipal disaster administration and field center operations through a three-tier architecture:\n"
        "(a) Presentation Tier: A responsive Single-Page Application (SPA) built using React 18, Vite, and Tailwind CSS. The interface provides accessible, high-contrast light and dark modes, intuitive tabular data tables, real-time KPI progress cards, camera-driven HTML5 QR scanning, and interactive analytics visualizations powered by Recharts.\n"
        "(b) Application & Business Logic Tier: A robust Laravel backend architected using Action-Domain-Responder (ADR) and Service-Repository patterns. Each functional capability is encapsulated in discrete Action classes (e.g., SearchHouseholdsAction, UpdateEvacuatedMemberStatusAction, CreateResourceRequestAction) with typed DTOs and Form Request validation.\n"
        "(c) Persistence Tier: A MySQL 8.0 relational database operating with InnoDB engine, utf8mb4 encoding, strict foreign key constraints, and soft deletion flags to maintain forensic audit trails.",
        bold_prefix="System Architecture: "
    )
    add_body_p(doc,
        "System Interoperability Context:\n"
        "EvaTrack interfaces with SafeTrack via private authenticated API channels to ingest household rosters and vulnerability markers. SafeTrack data remains confidential to system operators. Concurrently, EvaTrack communicates with ResQperation via outbound webhooks and REST endpoints to transfer resource requests and trigger push notification delivery to registered resident devices.",
        bold_prefix="Ecosystem Placement: "
    )

    add_styled_heading(doc, "2.2 Product Functions", level=2)
    add_body_p(doc, "At a macro level, EvaTrack fulfills the following primary functions:")
    add_bullet(doc, "Ingest and index master citizen records from SafeTrack, restricting visibility solely to verification personnel.", bold_prefix="Demographic Ingestion: ")
    add_bullet(doc, "Rapid admission of arriving families using camera-based QR code scanning or indexed multi-criteria search (name, contact, address).", bold_prefix="Household Verification: ")
    add_bullet(doc, "Real-time state tracking of households and individual family members as Evacuated or Not Evacuated.", bold_prefix="Status Tracking: ")
    add_bullet(doc, "Dynamic adjustment of center-level and room-level capacity counts immediately upon household admission or checkout.", bold_prefix="Capacity Recalculation: ")
    add_bullet(doc, "Multi-center operational control, room zoning, and designated personnel assignment.", bold_prefix="Shelter & Room Management: ")
    add_bullet(doc, "Logging of shelter shortages (food, water, medicines) and dispatching resource assistance tickets directly to ResQperation.", bold_prefix="Resource Demand Routing: ")
    add_bullet(doc, "Logging and lifecycle resolution tracking of shelter facility damage, sanitation issues, and medical emergencies.", bold_prefix="In-Center Incident Reporting: ")
    add_bullet(doc, "Multi-channel broadcast transmission (SMS via GSM gateway and push notifications via ResQperation).", bold_prefix="Emergency Alert Dispatch: ")
    add_bullet(doc, "Real-time computation of demographic vulnerabilities (age cohorts, pregnant women, PWDs, sex breakdown, total affected persons).", bold_prefix="Evacuation Analytics: ")
    add_bullet(doc, "Immutable audit logging of operator actions with emergency override capabilities for system recovery.", bold_prefix="Oversight & Audit: ")

    add_styled_heading(doc, "2.3 User Classes and Characteristics", level=2)
    user_table = doc.add_table(rows=7, cols=3)
    user_table.cell(0, 0).text = "User Class"
    user_table.cell(0, 1).text = "Operational Responsibilities & Scope"
    user_table.cell(0, 2).text = "Proficiency & Access Level"

    users_data = [
        ("Super Admin", 
         "Maintains supreme oversight across all centers, users, and disaster events. Generates system-wide analytics, exports consolidated municipality reports, and holds emergency override authority across all administrative actions for troubleshooting.", 
         "High technical proficiency; Full root system access."),
        ("Evacuation Admin", 
         "Creates and configures evacuation centers and sub-rooms. Assigns evacuation personnel to specific shelters. Monitors cross-center occupancy, active resource requests, and issue reports. Authorizes and dispatches emergency broadcasts via SMS and ResQperation push.", 
         "Moderate; Trained disaster operation manager."),
        ("Evacuation Personnel", 
         "Stationed on-site at assigned evacuation centers. Performs QR scanning and manual search verification of arriving evacuees. Updates household stay details, assigns families to available rooms, logs facility needs/shortages, and submits assistance requests.", 
         "Basic to moderate; Field staff trained in mobile/tablet scanner operation."),
        ("Households / Residents\n(Secondary Actor)", 
         "Displaced citizens who receive emergency broadcast warnings via SMS or ResQperation mobile app. Present physical or digital QR codes upon center arrival. Do not possess interactive login credentials.", 
         "General public; Zero technical training required."),
        ("SafeTrack Platform\n(External Actor)", 
         "Authoritative municipal demographic database. Supplies verified household data, member rosters, and vulnerability tags to EvaTrack via automated API synchronization.", 
         "Machine-to-machine authenticated REST client."),
        ("ResQperation System\n(External Actor)", 
         "External disaster response platform. Ingests escalated resource/personnel tickets from EvaTrack and routes mobile push broadcasts to citizens.", 
         "Machine-to-machine authenticated REST / Webhook client.")
    ]
    for r_idx, (role, desc, prof) in enumerate(users_data, start=1):
        user_table.cell(r_idx, 0).text = role
        user_table.cell(r_idx, 1).text = desc
        user_table.cell(r_idx, 2).text = prof

    style_table(user_table, col_widths=[Inches(1.5), Inches(3.6), Inches(1.4)])

    add_styled_heading(doc, "2.4 Operating Environment", level=2)
    add_bullet(doc, "Ubuntu 22.04 / 24.04 LTS, Nginx 1.24+, PHP 8.2+ with PHP-FPM, Composer 2.x, MySQL 8.0 Enterprise/Community Server.", bold_prefix="Server Infrastructure: ")
    add_bullet(doc, "Modern evergreen browsers (Google Chrome 115+, Microsoft Edge 115+, Mozilla Firefox 115+, Safari 16+) on desktop, laptop, tablet, or smartphone viewports (360px to 4K).", bold_prefix="Client Platforms: ")
    add_bullet(doc, "Built-in or USB webcams, smartphone cameras supporting MediaDevices.getUserMedia() for real-time QR scanning; standard office printers for PDF export printing.", bold_prefix="Hardware & Peripherals: ")
    add_bullet(doc, "Standard TCP/IP with minimum 256 kbps bandwidth for field personnel; HTTPS (TLS 1.2/1.3); cellular GSM connectivity for SMS dispatch gateways.", bold_prefix="Network Environment: ")

    add_styled_heading(doc, "2.5 Design and Implementation Constraints", level=2)
    add_bullet(doc, "The backend must strictly utilize Laravel 11/12 with PHP 8.2+ following Domain-Driven Design (DDD) with discrete Action, DTO, and Repository layers.", bold_prefix="Backend Architectural Constraint: ")
    add_bullet(doc, "The frontend must be implemented in React 18 with Vite build tooling and Tailwind CSS utility styling for accessible dark/light presentation.", bold_prefix="Frontend Architectural Constraint: ")
    add_bullet(doc, "Every query execution must utilize Eloquent ORM or parameterized query binding; string concatenation in SQL queries is prohibited.", bold_prefix="Database Security Constraint: ")
    add_bullet(doc, "In accordance with RA 10173, SafeTrack demographic data must never be published, displayed on public portals, or accessed by unauthorized roles.", bold_prefix="Privacy & Statutory Constraint: ")
    add_bullet(doc, "All database schema changes must be codified as reversible Laravel migrations version-controlled within Git.", bold_prefix="Version Control Constraint: ")

    add_styled_heading(doc, "2.6 Assumptions and Dependencies", level=2)
    add_bullet(doc, "Municipal LGUs designate official evacuation centers and appoint certified Evacuation Admins and Personnel prior to disaster onset.", bold_prefix="Operational Assumption: ")
    add_bullet(doc, "SafeTrack demographic registries are populated with accurate citizen rosters prior to disaster event declaration.", bold_prefix="Data Dependency: ")
    add_bullet(doc, "ResQperation server endpoints remain available to accept logistics requests and deliver push messages during active emergencies.", bold_prefix="Integration Dependency: ")
    add_bullet(doc, "Cellular telco carriers (Smart, Globe, DITO) maintain functional SMS gateway relay capability during evacuation advisories.", bold_prefix="Telecommunications Dependency: ")

    add_styled_heading(doc, "2.7 Apportioning of Requirements", level=2)
    add_body_p(doc,
        "All requirements cataloged under the 'Must' priority in Section 3 represent the non-negotiable core must-haves baseline specified in EVATRACK-Must-Haves (1).docx and verified for the panel defense. 'Should' priority requirements represent operational enhancements implemented to optimize usability and field resilience. 'Could' priority items represent exploratory extensions reserved for post-capstone production iterations."
    )

    doc.add_page_break()

    # =========================================================================
    # 3. SPECIFIC REQUIREMENTS
    # =========================================================================
    add_styled_heading(doc, "3. Specific Requirements", level=1)

    add_styled_heading(doc, "3.1 External Interface Requirements", level=2)

    add_styled_heading(doc, "3.1.1 User Interfaces", level=3)
    add_bullet(doc, "The system shall deliver a fully responsive, mobile-first web interface rendered via React and Tailwind CSS, supporting display resolutions from 360x640 up to 3840x2160.", bold_prefix="UI-01 Responsive Layout: ")
    add_bullet(doc, "The user interface shall support instantaneous, flash-free toggling between dark mode and light mode, persisting preference in localStorage.", bold_prefix="UI-02 Theme Adaptation: ")
    add_bullet(doc, "The dashboard shall display an uncluttered emergency announcement banner indicating the active disaster event name, type, and severity with manual pagination controls.", bold_prefix="UI-03 Emergency Announcement Banner: ")
    add_bullet(doc, "All tabular lists (Households, Evacuees, Centers, Issues, Requests) shall provide real-time search filtering, column sort toggles, status pills, and accessible pagination.", bold_prefix="UI-04 Tabular Interfaces: ")
    add_bullet(doc, "The verification workflow shall provide a high-performance camera viewfinder modal with real-time target framing, automatic camera switching, and immediate audio-visual scan feedback.", bold_prefix="UI-05 QR Scanner Viewfinder: ")

    add_styled_heading(doc, "3.1.2 Hardware Interfaces", level=3)
    add_bullet(doc, "Standard integrated or external USB/Bluetooth video capture devices adhering to WebRTC / MediaDevices APIs for optical QR code scanning.", bold_prefix="HW-01 Camera Interface: ")
    add_bullet(doc, "Standard operating system print subsystem (CUPS, Windows Spooler) accessible via window.print() for generating paper evacuation manifests and CSV logs.", bold_prefix="HW-02 Print Interface: ")

    add_styled_heading(doc, "3.1.3 Software Interfaces", level=3)
    sw_table = doc.add_table(rows=5, cols=3)
    sw_table.cell(0, 0).text = "Interface Name"
    sw_table.cell(0, 1).text = "Functional Scope & Purpose"
    sw_table.cell(0, 2).text = "Protocol & Integration Mechanism"

    sw_data = [
        ("SafeTrack Demographics API", 
         "Retrieves household profiles, family member rosters, addresses, and vulnerability categories (infants, elderly, pregnant, PWDs).", 
         "REST / JSON over HTTPS; Bearer API Token; Local database replication sync."),
        ("ResQperation Push Gateway", 
         "Transmits urgent disaster warnings and push notifications to registered residents holding the ResQperation mobile app.", 
         "Outbound REST API / JSON over HTTPS; Asynchronous queue worker."),
        ("ResQperation Logistics Dispatch", 
         "Escalates center supply demands (food, water, medicine, personnel) to ResQperation response coordinators.", 
         "REST API / Webhook payload containing itemized requisition DTOs."),
        ("Cellular SMS Gateway", 
         "Transmits SMS emergency evacuation alerts and verification acknowledgments to residents lacking internet connectivity.", 
         "HTTPS REST / Semaphore / Twilio SMS Gateway; Asynchronous job queue.")
    ]
    for r_idx, (name, scope, proto) in enumerate(sw_data, start=1):
        sw_table.cell(r_idx, 0).text = name
        sw_table.cell(r_idx, 1).text = scope
        sw_table.cell(r_idx, 2).text = proto

    style_table(sw_table, col_widths=[Inches(1.8), Inches(3.0), Inches(1.7)])

    add_styled_heading(doc, "3.1.4 Communications Interfaces", level=3)
    add_bullet(doc, "All browser-to-server and server-to-server exchanges shall use HTTPS with TLS 1.2/1.3 encryption.", bold_prefix="COM-01 Transport Encryption: ")
    add_bullet(doc, "API sessions shall authenticate via Laravel Sanctum bearer tokens with configurable expiration and automatic revocation upon logout.", bold_prefix="COM-02 Token Security: ")
    add_bullet(doc, "All state-modifying requests shall be protected against Cross-Site Request Forgery (CSRF) and Cross-Origin Resource Sharing (CORS) whitelisting.", bold_prefix="COM-03 Request Integrity: ")

    # =========================================================================
    # 3.2 FUNCTIONAL REQUIREMENTS
    # =========================================================================
    add_styled_heading(doc, "3.2 Functional Requirements", level=2)
    add_body_p(doc,
        "Functional requirements are cataloged below by operational domain. Each requirement is expressed as a verifiable 'shall' statement, assigned a unique identifier, categorized by MoSCoW priority (Must, Should, Could), and explicitly attributed to the responsible proponent and must-haves baseline from EVATRACK-Must-Haves (1).docx."
    )

    def add_fr_table(heading_text, reqs):
        add_styled_heading(doc, heading_text, level=3)
        t = doc.add_table(rows=len(reqs)+1, cols=4)
        t.cell(0, 0).text = "Req. ID"
        t.cell(0, 1).text = "Verifiable Requirement Description"
        t.cell(0, 2).text = "Priority"
        t.cell(0, 3).text = "Attribution"
        for r_idx, (r_id, desc, prio, attr) in enumerate(reqs, start=1):
            t.cell(r_idx, 0).text = r_id
            t.cell(r_idx, 1).text = desc
            t.cell(r_idx, 2).text = prio
            t.cell(r_idx, 3).text = attr
        style_table(t, col_widths=[Inches(1.1), Inches(3.6), Inches(0.7), Inches(1.1)])

    # 3.2.1 AUTH
    auth_reqs = [
        ("FR-AUTH-01", "The system shall authenticate operators using unique email/user ID credentials and bcrypt password hashes before granting access.", "Must", "Anna Rhea Villadolid"),
        ("FR-AUTH-02", "The system shall enforce Role-Based Access Control (RBAC) across Super Admin, Evacuation Admin, and Evacuation Personnel roles via backend route middleware.", "Must", "Anna Rhea Villadolid"),
        ("FR-AUTH-03", "The system shall restrict Evacuation Personnel actions strictly to their assigned evacuation center, denying modification of external centers.", "Must", "Anna Rhea Villadolid"),
        ("FR-AUTH-04", "The system shall allow Super Admins and Evacuation Admins to create, update, activate, and deactivate personnel accounts with center assignment.", "Must", "Anna Rhea Villadolid"),
        ("FR-AUTH-05", "The system shall invalidate user authorization tokens immediately upon logout and enforce session timeout after 30 minutes of inactivity.", "Must", "Anna Rhea Villadolid"),
        ("FR-AUTH-06", "The system shall allow authenticated operators to update their profile information and change passwords with current-password confirmation.", "Should", "Anna Rhea Villadolid")
    ]
    add_fr_table("3.2.1 Authentication & Role-Based Access Control (AUTH)", auth_reqs)

    # 3.2.2 STK
    stk_reqs = [
        ("FR-STK-01", "The system shall ingest household demographic data from SafeTrack, including head of household, member roster, birthdates, gender, and address.", "Must", "Vhenz Cernal"),
        ("FR-STK-02", "The system shall capture vulnerable group classifications from SafeTrack (infants, children, elderly, pregnant women, and PWDs).", "Must", "Vhenz Cernal"),
        ("FR-STK-03", "The system shall strictly restrict the visibility of SafeTrack demographic data to authorized evacuation personnel and admins during verification.", "Must", "Vhenz Cernal"),
        ("FR-STK-04", "The system shall prohibit direct public and resident access to raw SafeTrack demographic records to ensure compliance with RA 10173.", "Must", "Vhenz Cernal"),
        ("FR-STK-05", "The system shall support incremental re-synchronization of SafeTrack demographic datasets without overwriting active evacuation stay logs.", "Should", "Vhenz Cernal")
    ]
    add_fr_table("3.2.2 SafeTrack Demographics Ingestion & Isolation (STK)", stk_reqs)

    # 3.2.3 CTR
    ctr_reqs = [
        ("FR-CTR-01", "The system shall allow Evacuation Admins and Super Admins to create and configure evacuation centers with name, address coordinates, and capacity.", "Must", "Vhenz Cernal"),
        ("FR-CTR-02", "The system shall allow Evacuation Admins to configure individual rooms and partitions within an evacuation center, specifying room capacity and gender allocation.", "Must", "Danica Gelbolingo"),
        ("FR-CTR-03", "The system shall automatically recalculate total center capacity and room occupancy upon household verification and room assignment.", "Must", "Vhenz Cernal"),
        ("FR-CTR-04", "The system shall display real-time occupancy percentage badges and progress meters indicating available, near-capacity, and full shelter status.", "Must", "Vhenz Cernal"),
        ("FR-CTR-05", "The system shall allow Evacuation Admins to assign and re-assign Evacuation Personnel to operational evacuation centers.", "Must", "Danica Gelbolingo"),
        ("FR-CTR-06", "The system shall prevent admitting households to centers that have reached 100% capacity unless an administrative override is authorized.", "Should", "Vhenz Cernal")
    ]
    add_fr_table("3.2.3 Evacuation Center & Room Configuration (CTR)", ctr_reqs)

    # 3.2.4 VER
    ver_reqs = [
        ("FR-VER-01", "The system shall enable Evacuation Personnel to verify arriving households via camera-driven optical scanning of digital or printed QR codes.", "Must", "Klint M. Ruales"),
        ("FR-VER-02", "The system shall provide fallback manual verification by searching SafeTrack records using household head name, member name, or contact number.", "Must", "Klint M. Ruales"),
        ("FR-VER-03", "The system shall display the household member roster upon successful verification, allowing personnel to select which members are physically present.", "Must", "Klint M. Ruales"),
        ("FR-VER-04", "The system shall update household evacuation status to 'Evacuated' once verified members are admitted to the center.", "Must", "Vhenz Cernal"),
        ("FR-VER-05", "The system shall automatically lock the 'Admit' button and display an 'Admitted' badge when a household reaches 100% verified member admission.", "Must", "Klint M. Ruales"),
        ("FR-VER-06", "The system shall record verification timestamps, admitting personnel ID, and admission mode (QR vs. Manual) in the persistent evacuation record.", "Must", "Klint M. Ruales"),
        ("FR-VER-07", "The system shall alert personnel if a household has already been admitted to another center to prevent fraudulent duplicate registrations.", "Must", "Klint M. Ruales")
    ]
    add_fr_table("3.2.4 Household Verification & Admission Tracking (VER)", ver_reqs)

    # 3.2.5 STAY
    stay_reqs = [
        ("FR-STAY-01", "The system shall allow Evacuation Personnel to assign admitted households and family members to designated rooms within the center.", "Must", "Danica Gelbolingo"),
        ("FR-STAY-02", "The system shall monitor household stay duration, displaying admission timestamp and real-time elapsed hours/days in shelter.", "Must", "Danica Gelbolingo"),
        ("FR-STAY-03", "The system shall allow personnel to update individual member statuses (e.g., in-shelter, referred to hospital, temporary leave).", "Must", "Danica Gelbolingo"),
        ("FR-STAY-04", "The system shall allow Evacuation Personnel to process household checkout upon official evacuation decampment.", "Must", "Danica Gelbolingo"),
        ("FR-STAY-05", "The system shall automatically deduct departing household member counts from center and room occupancy tallies upon checkout.", "Must", "Vhenz Cernal"),
        ("FR-STAY-06", "The system shall archive completed evacuation records permanently for historical analytics while updating active status to 'Decamped'.", "Must", "Danica Gelbolingo")
    ]
    add_fr_table("3.2.5 In-Center Stay, Room Assignment & Checkout (STAY)", stay_reqs)

    # 3.2.6 EVT
    evt_reqs = [
        ("FR-EVT-01", "The system shall allow Evacuation Admins and Super Admins to declare and manage disaster events with event name, disaster type, and severity.", "Must", "Vhenz Cernal"),
        ("FR-EVT-02", "The system shall automatically link operational evacuation centers to active disaster declarations across the municipality.", "Must", "Vhenz Cernal"),
        ("FR-EVT-03", "The system shall display active disaster announcements prominently on the operator dashboard without distracting marquee animation.", "Must", "Klint M. Ruales"),
        ("FR-EVT-04", "The system shall allow Evacuation Admins to conclude active disaster events, archiving event-scoped evacuation metrics.", "Must", "Vhenz Cernal"),
        ("FR-EVT-05", "The system shall track historical disaster events, allowing comparative analytics across past typhoons, floods, and earthquakes.", "Should", "Klint M. Ruales")
    ]
    add_fr_table("3.2.6 Disaster Event Coordination (EVT)", evt_reqs)

    # 3.2.7 ISS
    iss_reqs = [
        ("FR-ISS-01", "The system shall allow Evacuation Personnel to log in-center issue reports categorized by Facility Damage, Sanitation, Health, or Security.", "Must", "Anna Rhea Villadolid"),
        ("FR-ISS-02", "The system shall capture issue severity levels (Low, Medium, High, Critical) and incident descriptions for administrative triage.", "Must", "Anna Rhea Villadolid"),
        ("FR-ISS-03", "The system shall enable Evacuation Admins to review, acknowledge, assign, and resolve center issue reports.", "Must", "Anna Rhea Villadolid"),
        ("FR-ISS-04", "The system shall maintain an audit history of issue status transitions (Open -> In Progress -> Resolved -> Closed).", "Must", "Anna Rhea Villadolid"),
        ("FR-ISS-05", "The system shall display center issue distribution charts and resolution rate KPIs on the administrative analytics dashboard.", "Must", "Klint M. Ruales")
    ]
    add_fr_table("3.2.7 In-Center Needs & Incident Reporting (ISS)", iss_reqs)

    # 3.2.8 RES
    res_reqs = [
        ("FR-RES-01", "The system shall allow Evacuation Personnel to submit resource supply requests specifying resource type, quantity demanded, and urgency.", "Must", "Anna Rhea Villadolid"),
        ("FR-RES-02", "The system shall allow Evacuation Personnel to request additional response personnel and specialized rescue support from ResQperation.", "Must", "Anna Rhea Villadolid"),
        ("FR-RES-03", "The system shall route resource requests to Evacuation Admins and transmit external assistance payloads to ResQperation coordinators.", "Must", "Anna Rhea Villadolid"),
        ("FR-RES-04", "The system shall track resource fulfillment lifecycle across Pending, Approved, In Transit, Delivered, and Rejected states.", "Must", "Anna Rhea Villadolid"),
        ("FR-RES-05", "The system shall visualize top requested supply categories and volumes in a proportional 3-column progress chart with right-aligned units.", "Must", "Klint M. Ruales"),
        ("FR-RES-06", "The system shall compute overall resource fulfillment rates per center and system-wide.", "Must", "Klint M. Ruales")
    ]
    add_fr_table("3.2.8 Resource Demands & ResQperation Response Routing (RES)", res_reqs)

    # 3.2.9 NOTIF
    notif_reqs = [
        ("FR-NOTIF-01", "The system shall provide an SMS & Push dispatch interface for Evacuation Admins to compose and transmit urgent public evacuation advisories.", "Must", "Klint M. Ruales"),
        ("FR-NOTIF-02", "The system shall route SMS alerts to registered household contact numbers lacking active mobile data via cellular SMS gateway.", "Must", "Klint M. Ruales"),
        ("FR-NOTIF-03", "The system shall route push notification advisories to residents holding the ResQperation mobile application.", "Must", "Klint M. Ruales"),
        ("FR-NOTIF-04", "The system shall allow geographic targeting of advisories by municipality, barangay, or specific evacuation zone.", "Must", "Klint M. Ruales"),
        ("FR-NOTIF-05", "The system shall track transmission dispatch logs, recipient counts, and delivery success/failure statuses.", "Must", "Klint M. Ruales"),
        ("FR-NOTIF-06", "The system shall export broadcast history logs to CSV for post-disaster audit reporting.", "Should", "Klint M. Ruales")
    ]
    add_fr_table("3.2.9 Emergency Dispatches, SMS & Push Broadcasts (NOTIF)", notif_reqs)

    # 3.2.10 ANL
    anl_reqs = [
        ("FR-ANL-01", "The system shall dynamically aggregate evacuation demographic analytics from verified admission records.", "Must", "Klint M. Ruales"),
        ("FR-ANL-02", "The system shall compute age group distributions (infants 0-2, children 3-12, adolescents 13-17, adults 18-59, elderly 60+).", "Must", "Klint M. Ruales"),
        ("FR-ANL-03", "The system shall compute total counts and percentages of pregnant women and nursing mothers admitted to evacuation centers.", "Must", "Klint M. Ruales"),
        ("FR-ANL-04", "The system shall compute total counts of Persons with Disabilities (PWDs) and individuals requiring critical medical assistance.", "Must", "Klint M. Ruales"),
        ("FR-ANL-05", "The system shall compute gender distribution tallies (male vs. female) across verified evacuees.", "Must", "Klint M. Ruales"),
        ("FR-ANL-06", "The system shall compute total affected population, total evacuated households, and municipal shelter capacity utilization rates.", "Must", "Klint M. Ruales"),
        ("FR-ANL-07", "The system shall filter analytics by Evacuation Center and Date Range and export summary reports to PDF and CSV formats.", "Must", "Klint M. Ruales")
    ]
    add_fr_table("3.2.10 Evacuation Analytics & Demographic Reporting (ANL)", anl_reqs)

    # 3.2.11 AUD
    aud_reqs = [
        ("FR-AUD-01", "The system shall record immutable audit log entries (operator ID, timestamp, IP address, action, target entity) for all security and data operations.", "Must", "Klint M. Ruales"),
        ("FR-AUD-02", "The system shall grant Super Admins full oversight of system operations across all evacuation centers and administrative roles.", "Must", "Klint M. Ruales"),
        ("FR-AUD-03", "The system shall provide Super Admins with emergency override capabilities to intervene, correct, or unlock records during field emergencies.", "Must", "Klint M. Ruales"),
        ("FR-AUD-04", "The system shall log every administrative override action with mandatory justification notes to ensure accountability.", "Must", "Klint M. Ruales"),
        ("FR-AUD-05", "The system shall generate consolidated, municipality-wide evacuation summary reports exportable by Super Admins.", "Must", "Klint M. Ruales")
    ]
    add_fr_table("3.2.11 Audit Logging & Super Admin System Oversight (AUD)", aud_reqs)

    doc.add_page_break()

    # =========================================================================
    # 3.3 NON-FUNCTIONAL REQUIREMENTS
    # =========================================================================
    add_styled_heading(doc, "3.3 Non-Functional Requirements", level=2)

    add_styled_heading(doc, "3.3.1 Performance", level=3)
    add_bullet(doc, "Under a concurrent load of 50 field personnel across 10 centers, 95% of API requests shall complete within 1.5 seconds.", bold_prefix="NFR-PER-01 Response Time: ")
    add_bullet(doc, "Optical QR code detection and household lookup shall execute and display results in under 500 milliseconds on standard smartphone hardware.", bold_prefix="NFR-PER-02 Scanning Latency: ")
    add_bullet(doc, "Full-text household search over 20,000 SafeTrack citizen records shall return paginated matches within 1.0 second.", bold_prefix="NFR-PER-03 Search Throughput: ")
    add_bullet(doc, "Analytics KPI computations over 100,000 historic member records shall execute in under 3.0 seconds.", bold_prefix="NFR-PER-04 Analytics Aggregation: ")
    add_bullet(doc, "The frontend production bundle shall achieve a Lighthouse Performance score >= 85 and load within 2 seconds over 4G connections.", bold_prefix="NFR-PER-05 Frontend Efficiency: ")

    add_styled_heading(doc, "3.3.2 Security & Data Confidentiality", level=3)
    add_bullet(doc, "All operator passwords must be hashed using bcrypt with a work factor of 12; plaintext passwords shall never be persisted or logged.", bold_prefix="NFR-SEC-01 Cryptographic Storage: ")
    add_bullet(doc, "All database queries must use PDO prepared statements via Eloquent ORM to eliminate SQL injection vulnerabilities.", bold_prefix="NFR-SEC-02 Parameterized Execution: ")
    add_bullet(doc, "State-modifying endpoints must enforce CSRF token verification and strict Content Security Policy (CSP) headers.", bold_prefix="NFR-SEC-03 Web Application Armor: ")
    add_bullet(doc, "SafeTrack demographic records must be isolated behind authenticated server-side controllers and never exposed via public endpoints.", bold_prefix="NFR-SEC-04 Demographic Isolation: ")
    add_bullet(doc, "The system shall enforce account lockout for 15 minutes after 5 consecutive failed login attempts.", bold_prefix="NFR-SEC-05 Brute-Force Throttling: ")
    add_bullet(doc, "Authorization must be evaluated on the server per request; hiding frontend buttons shall never serve as sole security enforcement.", bold_prefix="NFR-SEC-06 Server-Side Enforcement: ")

    add_styled_heading(doc, "3.3.3 Reliability and Disaster Availability", level=3)
    add_bullet(doc, "The system shall maintain 99.5% operational availability during active disaster declaration windows.", bold_prefix="NFR-REL-01 Emergency Uptime: ")
    add_bullet(doc, "Automated database backups must execute daily with retention for 30 days; point-in-time recovery must be verified within 30 minutes.", bold_prefix="NFR-REL-02 Backup & Recovery: ")
    add_bullet(doc, "Database transactions must satisfy strict ACID properties; failed admissions shall roll back completely to prevent orphan records.", bold_prefix="NFR-REL-03 Transactional Integrity: ")
    add_bullet(doc, "Field personnel must be able to switch instantaneously to manual search verification if camera hardware or QR printouts fail.", bold_prefix="NFR-REL-04 Operational Fallback: ")

    add_styled_heading(doc, "3.3.4 Usability & Field Accessibility", level=3)
    add_bullet(doc, "Evacuation personnel with basic digital literacy shall complete a household admission in under 45 seconds via QR code.", bold_prefix="NFR-USA-01 Admission Velocity: ")
    add_bullet(doc, "The UI must satisfy WCAG 2.1 Level AA contrast ratios (minimum 4.5:1 for standard text) across both dark and light modes.", bold_prefix="NFR-USA-02 Visual Contrast: ")
    add_bullet(doc, "All interactive buttons and touch targets must measure at least 44x44 CSS pixels for accurate field operation on tablets.", bold_prefix="NFR-USA-03 Touch Usability: ")
    add_bullet(doc, "Validation errors and system alerts must provide descriptive, non-technical instructions guiding corrective action.", bold_prefix="NFR-USA-04 Error Clarity: ")

    add_styled_heading(doc, "3.3.5 Maintainability & Clean Architecture", level=3)
    add_bullet(doc, "The backend must adhere to PSR-12 coding style and Domain-Driven Design (DDD) separating Actions, Repositories, and Models.", bold_prefix="NFR-MNT-01 Architecture Standards: ")
    add_bullet(doc, "Automated PHPUnit/Pest feature tests must cover 100% of Must-priority functional workflows prior to release.", bold_prefix="NFR-MNT-02 Automated Verification: ")
    add_bullet(doc, "Every database mutation must be codified as version-controlled, rollback-capable Laravel migration files.", bold_prefix="NFR-MNT-03 Migration Reproducibility: ")
    add_bullet(doc, "Frontend React components must be modular, adhering to single responsibility with typed store abstractions.", bold_prefix="NFR-MNT-04 Component Modularity: ")

    add_styled_heading(doc, "3.3.6 Portability", level=3)
    add_bullet(doc, "The entire application must be deployable across standard cloud VPS or on-premises municipal servers without code alteration.", bold_prefix="NFR-POR-01 Environment Agnostic: ")
    add_bullet(doc, "The application must support multi-container Docker deployment (Nginx, PHP-FPM, MySQL) for seamless reproduction.", bold_prefix="NFR-POR-02 Containerization: ")
    add_bullet(doc, "The frontend client must function flawlessly across Google Chrome, Mozilla Firefox, Microsoft Edge, and Apple Safari.", bold_prefix="NFR-POR-03 Browser Compatibility: ")

    add_styled_heading(doc, "3.3.7 Data Privacy and Regulatory Compliance", level=3)
    add_bullet(doc, "Full adherence to Republic Act No. 10173 (Data Privacy Act of 2012) principles of transparency, legitimate purpose, and proportionality.", bold_prefix="NFR-CMP-01 RA 10173 Compliance: ")
    add_bullet(doc, "SafeTrack demographic records must be utilized exclusively for evacuation verification and needs planning, never commercialized.", bold_prefix="NFR-CMP-02 Purpose Limitation: ")
    add_bullet(doc, "Demographic analytics exports must be aggregated and anonymized to protect individual citizen identities in public reports.", bold_prefix="NFR-CMP-03 Anonymized Reporting: ")
    add_bullet(doc, "Adherence to NDRRMC Joint Memorandum Circular No. 1 guidelines for evacuation center registration and family tracking.", bold_prefix="NFR-CMP-04 NDRRMC Guidelines: ")

    add_styled_heading(doc, "3.4 Database Requirements", level=2)
    add_body_p(doc,
        "The EvaTrack relational persistence layer is implemented in MySQL 8.0 with InnoDB storage engine and utf8mb4_unicode_ci collation. The schema encompasses 35 dedicated entities designed around Domain-Driven boundaries. The core entities and their relationships are cataloged below:"
    )

    db_table = doc.add_table(rows=16, cols=3)
    db_table.cell(0, 0).text = "Entity Name"
    db_table.cell(0, 1).text = "Domain Scope & Description"
    db_table.cell(0, 2).text = "Key Cardinalities & Relationships"

    db_entities = [
        ("users", "System operators across Super Admin, Evacuation Admin, and Evacuation Personnel roles.", "N:1 roles; N:1 evacuation_centers (assigned center)"),
        ("roles", "System privilege classifications (super_admin, evac_admin, evac_personnel).", "1:N users"),
        ("households", "Master family unit ingested from SafeTrack with address and contact details.", "1:N household_members; 1:N evacuation_records; 1:N notification_recipients"),
        ("household_members", "Individual citizen records comprising birthdate, gender, and relationships.", "N:1 households; N:1 genders; 1:N member_vulnerable_groups; 1:N evacuated_members"),
        ("vulnerable_groups", "Classifications of vulnerability: Senior Citizen, PWD, Pregnant, Infant, Child.", "1:N member_vulnerable_groups"),
        ("member_vulnerable_groups", "Pivot entity linking individual members to one or more vulnerability tags.", "N:1 household_members; N:1 vulnerable_groups"),
        ("evacuation_centers", "Registered municipal shelter facilities with capacity and coordinates.", "1:N rooms; 1:N evacuation_records; 1:N users (personnel); 1:N resource_requests"),
        ("rooms", "Designated rooms or sub-zones within an evacuation center with individual limits.", "N:1 evacuation_centers; 1:N evacuated_members"),
        ("disaster_events", "Declared municipal disaster incidents with severity and operational lifecycle.", "N:1 disaster_types; N:1 severity_levels; 1:N evacuation_records"),
        ("evacuation_records", "Transactional record anchoring a household admission to a center and event.", "N:1 households; N:1 evacuation_centers; N:1 disaster_events; 1:N evacuated_members"),
        ("evacuated_members", "Tracks individual member check-in timestamp, assigned room, and status.", "N:1 evacuation_records; N:1 household_members; N:1 rooms; N:1 evacuation_statuses"),
        ("resource_requests", "Logistics tickets requested by center personnel and routed to ResQperation.", "N:1 evacuation_centers; N:1 resource_request_statuses; N:1 urgency_levels"),
        ("center_issue_reports", "In-center facility, medical, sanitation, and security incident reports.", "N:1 evacuation_centers; N:1 issue_statuses; N:1 severity_levels; N:1 categories"),
        ("notifications", "Broadcast transmission records dispatched via SMS and ResQperation push.", "1:N notification_recipients; 1:N notification_logs"),
        ("audit_logs", "Immutable forensic trace capturing operator actions, IP addresses, and before/after states.", "N:1 users")
    ]
    for r_idx, (ent, desc, rel) in enumerate(db_entities, start=1):
        db_table.cell(r_idx, 0).text = ent
        db_table.cell(r_idx, 1).text = desc
        db_table.cell(r_idx, 2).text = rel

    style_table(db_table, col_widths=[Inches(1.8), Inches(2.9), Inches(1.8)])

    add_body_p(doc,
        "Integrity Rules: Foreign key constraints enforce referential integrity with ON DELETE RESTRICT on all operational records to prevent accidental cascading destruction of active shelter data. Soft deletion (deleted_at timestamps) is utilized across users, households, centers, and records to preserve forensic trails.",
        bold_prefix="Data Integrity & Archival: "
    )

    doc.add_page_break()

    # =========================================================================
    # APPENDIX A. PRELIMINARY ER OVERVIEW
    # =========================================================================
    add_styled_heading(doc, "Appendix A. Preliminary Entity-Relationship Overview", level=1)
    add_body_p(doc,
        "The conceptual architecture of EvaTrack centers upon the evacuation_records transaction, which serves as the core operational nexus linking a SafeTrack household to an evacuation_center under an active disaster_event. The individual evacuated_members records capture fine-grained presence, room allocation, and decampment status. Supporting operational subsystems branch outward from evacuation_centers: rooms govern spatial allocations; resource_requests track logistics demands routed to ResQperation; center_issue_reports capture facility health; and notifications record multi-channel broadcasts. The full 35-entity diagram is cataloged in the software design artifact evatrack_erd_reference.md."
    )

    # =========================================================================
    # APPENDIX B. REQUIREMENTS TRACEABILITY MATRIX
    # =========================================================================
    add_styled_heading(doc, "Appendix B. Requirements Traceability Matrix", level=1)
    add_body_p(doc,
        "The following matrix maps every explicit Must-Have requirement from EVATRACK-Must-Haves (1).docx to its corresponding Functional Requirement ID, implemented backend/frontend modules, and designated student proponent:"
    )

    trace_table = doc.add_table(rows=18, cols=4)
    trace_table.cell(0, 0).text = "Baseline Must-Have Requirement"
    trace_table.cell(0, 1).text = "FR Mapping"
    trace_table.cell(0, 2).text = "Implementing Module / Action"
    trace_table.cell(0, 3).text = "Proponent"

    trace_data = [
        ("Retrieve SafeTrack demographic data for identification & verification only; data hidden from households.", "FR-STK-01\nFR-STK-03\nFR-STK-04", "SafeTrackSyncService\nHouseholdController", "Vhenz Cernal"),
        ("Track household evacuation status as Evacuated or Not Evacuated.", "FR-VER-04\nFR-STAY-03", "UpdateEvacuatedMemberStatusAction\nHouseholdRepository", "Vhenz Cernal"),
        ("Automatically update evacuation center capacity when household is verified.", "FR-CTR-03\nFR-VER-04", "AdmitHouseholdAction\nCenterCapacityService", "Vhenz Cernal"),
        ("Compute evacuation analytics from demographics (age, pregnant, PWD, gender, affected pop).", "FR-ANL-01\nFR-ANL-02\nFR-ANL-06", "LiveAnalyticsService\nAnalyticsKPIs.jsx", "Klint M. Ruales"),
        ("Route alerts and assistance requests to households and ResQperation.", "FR-RES-03\nFR-NOTIF-02\nFR-NOTIF-03", "ResQperationDispatchService\nSMSGatewayService", "Anna Rhea Villadolid"),
        ("Evacuation Personnel: Log in (evacuation personnel role).", "FR-AUTH-01\nFR-AUTH-02", "LoginAction\nRoleMiddleware", "Anna Rhea Villadolid"),
        ("Verify evacuated households via QR code scanning or manual input.", "FR-VER-01\nFR-VER-02\nFR-VER-03", "SearchHouseholdsAction\nRegistrySearch.jsx", "Klint M. Ruales"),
        ("Update household info post-verification: assign rooms, monitor stay, update status.", "FR-STAY-01\nFR-STAY-02\nFR-STAY-03", "AssignRoomAction\nEvacuationDetail.jsx", "Danica Gelbolingo"),
        ("Report in-center needs (shortages of resources and extra personnel requests).", "FR-ISS-01\nFR-RES-01\nFR-RES-02", "CreateCenterIssueReportAction\nCreateResourceRequestAction", "Anna Rhea Villadolid"),
        ("Request resources and personnel assistance from ResQperation.", "FR-RES-02\nFR-RES-03", "ResQperationRequestService\nResourceRequests.jsx", "Anna Rhea Villadolid"),
        ("Evacuation Admin: Log in (evacuation admin role).", "FR-AUTH-01\nFR-AUTH-02", "LoginAction\nRoleMiddleware", "Anna Rhea Villadolid"),
        ("Create and configure evacuation centers.", "FR-CTR-01\nFR-CTR-04", "CreateCenterAction\nEvacuationList.jsx", "Vhenz Cernal"),
        ("Create and configure rooms within an evacuation center.", "FR-CTR-02", "CreateRoomAction\nRoomManagementModal.jsx", "Danica Gelbolingo"),
        ("Monitor ongoing evacuations across centers (status, occupancy, personnel activity).", "FR-CTR-04\nFR-EVT-02", "LiveDashboardService\nDashboard.jsx", "Klint M. Ruales"),
        ("Monitor ongoing requests per evacuation centers.", "FR-RES-04\nFR-RES-06", "ListResourceRequestsAction\nResourceRequestsAnalytics.jsx", "Anna Rhea Villadolid"),
        ("Send evacuation alerts (SMS for non-internet; push via ResQperation for app).", "FR-NOTIF-01\nFR-NOTIF-02\nFR-NOTIF-03", "DispatchNotificationAction\nEvacuationAlerts.jsx", "Klint M. Ruales"),
        ("Super Admin: System-wide reports, full oversight, and emergency override access.", "FR-AUD-02\nFR-AUD-03\nFR-AUD-05", "AnalyticsExportController\nSuperAdminOverrideMiddleware", "Klint M. Ruales")
    ]
    for r_idx, (must, fr, mod, prop) in enumerate(trace_data, start=1):
        trace_table.cell(r_idx, 0).text = must
        trace_table.cell(r_idx, 1).text = fr
        trace_table.cell(r_idx, 2).text = mod
        trace_table.cell(r_idx, 3).text = prop

    style_table(trace_table, col_widths=[Inches(2.5), Inches(1.1), Inches(1.7), Inches(1.2)])

    doc.add_page_break()

    # =========================================================================
    # APPENDIX C. APPROVAL & SIGN-OFF SHEET
    # =========================================================================
    add_styled_heading(doc, "Appendix C. Capstone Approval & Sign-Off Sheet", level=1)
    add_body_p(doc,
        "By signing below, the capstone adviser, project coordinator, panel evaluators, and student proponents attest that this Software Requirements Specification has been thoroughly evaluated and approved as the definitive technical and functional baseline for EVATRACK (Version 1.0)."
    )

    sign_table = doc.add_table(rows=9, cols=3)
    sign_table.cell(0, 0).text = "Role in Capstone Project"
    sign_table.cell(0, 1).text = "Full Name & Academic Title"
    sign_table.cell(0, 2).text = "Signature / Date"

    signatories = [
        ("Student Proponent (Lead)", "Klint M. Ruales", "_________________________"),
        ("Student Proponent", "Danica Gelbolingo", "_________________________"),
        ("Student Proponent", "Anna Rhea Villadolid", "_________________________"),
        ("Student Proponent", "Vhenz Cernal", "_________________________"),
        ("Capstone Project Adviser", "Mr. Stefan A.D. Flores", "_________________________"),
        ("Capstone Project Coordinator", "Mr. Roderick A. Bandalan", "_________________________"),
        ("Panel Evaluation Chair", "Mr. Eric Magto", "_________________________"),
        ("Panel Evaluator", "Ms. Khiara Danielle Rubia", "_________________________")
    ]
    for r_idx, (role, name, sig) in enumerate(signatories, start=1):
        sign_table.cell(r_idx, 0).text = role
        sign_table.cell(r_idx, 1).text = name
        sign_table.cell(r_idx, 2).text = sig

    style_table(sign_table, col_widths=[Inches(2.2), Inches(2.3), Inches(2.0)])

    # Save Document
    doc.save(output_path)
    print(f"Successfully generated complete SRS at: {output_path}")

if __name__ == "__main__":
    out_docx = "c:/CAPSTONE/evatrack/docs/EVATRACK-SRS-v1.0.docx"
    build_srs_document(out_docx)
